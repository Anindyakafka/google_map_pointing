"""Build map import files from saved Google Maps candidates and explicit selections."""
import csv
import hashlib
import json
import math
import re
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from geocode import ROOT, source_rows


def normalized(value):
    return re.sub(r'[^a-z0-9]', '', value.lower())


def export():
    records = json.loads((ROOT / 'data/candidates.json').read_text(encoding='utf-8'))
    review_path = ROOT / 'data/selections.json'
    selections = json.loads(review_path.read_text(encoding='utf-8')) if review_path.exists() else {}
    extra_path = ROOT / 'data/extra_candidates.json'
    extra = json.loads(extra_path.read_text(encoding='utf-8')) if extra_path.exists() else {}
    source = source_rows()
    assert len(records) == len(source), 'Candidate collection is incomplete.'
    rows = []
    for original, record in zip(source, records):
        assert original == record['row'], 'Source workbook changed: recollect candidates.'
        row_id = str(original['Sr. No.'])
        decision = selections.get(row_id)
        searches = record['searches'] + extra.get(row_id, [])
        unique = {}
        for result in searches:
            for candidate in result['candidates']:
                unique.setdefault(candidate['google_id'], (candidate, result))
        chosen = None
        note = ''
        status = 'unresolved'
        if decision:
            note = decision['note']
            if decision.get('google_id'):
                chosen = unique[decision['google_id']]
                status = decision.get('status', 'reviewed_match')
        else:
            # Require agreement between district-only and block-qualified searches.
            matches = [pair for pair in unique.values() if normalized(pair[0]['name']) == normalized(original['Panchayat'])]
            if len(matches) == 1 and all(any(c['google_id'] == matches[0][0]['google_id'] for c in s['candidates']) for s in searches):
                chosen = matches[0]
                status = 'name_and_query_agreement'
                note = 'Exact place-name match returned for both block-qualified and district-qualified searches; locality point, not a GP boundary or verified office.'
        row = {key: original[key] for key in ('District', 'Block', 'Panchayat')}
        row.update(lat=None, long=None, Remarks=original['Remarks'])
        row.update({key: original[key] for key in ('Sr. No.', 'Generated', 'Comments')})
        row.update(match_status=status, matched_name='', matched_address='', coordinate_type='locality/place point', match_note=note,
                   google_maps_url='', google_id='', source_query='', source_url='', retrieved_utc='')
        if decision and decision.get('coordinate_type'):
            row['coordinate_type'] = decision['coordinate_type']
        if not chosen:
            row['coordinate_type'] = ''
        if chosen:
            candidate, result = chosen
            lat, lon = candidate['lat'], candidate['long']
            assert all(math.isfinite(v) for v in (lat, lon))
            assert 17.5 <= lat <= 22.7 and 81.3 <= lon <= 87.6, f'Outside Odisha review envelope: {row_id}'
            row.update(lat=round(lat, 7), long=round(lon, 7), matched_name=candidate['name'], matched_address=candidate['address'],
                       google_maps_url='https://www.google.com/maps?cid=' + str(int(candidate['google_id'].split(':')[1], 16)), google_id=candidate['google_id'],
                       source_query=result['query'], source_url=result['source_url'], retrieved_utc=result['retrieved_utc'])
        rows.append(row)
    output = ROOT / 'output'
    output.mkdir(exist_ok=True)
    fields = list(rows[0])
    for filename, subset in [('panchayats_coordinates.csv', rows),
                             ('panchayats_map_ready.csv', [r for r in rows if r['lat'] is not None and r['match_status'] != 'needs_review']),
                             ('panchayats_needs_review.csv', [r for r in rows if r['match_status'] in ('unresolved', 'needs_review')])]:
        with (output / filename).open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=fields)
            writer.writeheader()
            writer.writerows(subset)
    book = Workbook()
    sheet = book.active
    sheet.title = 'Sheet1'
    sheet.append(fields)
    for row in rows:
        sheet.append([row[field] for field in fields])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='245B78')
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(65, max(14, max(len(str(c.value or '')) for c in column) + 2))
    for row in sheet.iter_rows(min_row=2):
        row[3].number_format = row[4].number_format = '0.0000000'
        if row[9].value in ('unresolved', 'needs_review'):
            for cell in row:
                cell.fill = PatternFill('solid', fgColor='FFF2CC')
    book.save(output / 'panchayats_coordinates.xlsx')
    report = {'source_sheet': 'Sheet1', 'source_rows': len(source), 'output_rows': len(rows),
              'source_sha256': hashlib.sha256((ROOT / 'Odisha List.xlsx').read_bytes()).hexdigest(),
              'with_coordinates': sum(r['lat'] is not None for r in rows),
              'status_counts': dict(Counter(r['match_status'] for r in rows)),
              'district_counts': dict(Counter(r['District'] for r in rows))}
    (output / 'validation_summary.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    export()
